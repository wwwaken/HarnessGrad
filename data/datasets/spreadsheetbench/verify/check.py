#!/usr/bin/env python3
r"""The SpreadsheetBench check: recalculate the agent's workbook and compare it.

The comparison is the official one, transcribed from
`RUCKBReasoning/SpreadsheetBench`'s `evaluation/evaluation.py`:

    load both workbooks with `openpyxl.load_workbook(..., data_only=True)`
    for each `Sheet!A1:B2` range in the task's `answer_position`, compare every cell
    as `transform_value(gt) == transform_value(proc)`

`transform_value` (verbatim):

    numbers -> `round(float(v), 2)`
    `datetime` -> Excel's serial day count, rounded
    `datetime.time` -> `str(v)[:-3]` (to the minute)
    a string that parses as a float -> that float
    `""` and `None` compare equal

**Why the recalculation step exists.** `data_only=True` returns the value Excel last *cached*
next to each formula. A workbook produced by a program that computed its own answers has those
caches; one whose cells hold formulas that were never evaluated has `None` in every one of them,
and would be scored 0 for a reason no evidence would show. So the agent's workbook is
recalculated first, with LibreOffice in headless mode -- the same thing the repository's own
`evaluation/open_spreadsheet.py` does, and the reason `soffice` has to be installed.

**What is not compared.** Cell fill and font colour: the official `cell_level_compare` has both
comparisons present and commented out, and this file matches the code that runs.

**A check that cannot measure** -- unreadable workbooks, comparator/runtime failures, no
LibreOffice when the workbook holds formulas -- writes **no reward file**, writes evidence
saying `measured: false`, and exits 3. A bare 0 is indistinguishable from a wrong workbook.
An absent result artifact and a completed value/sheet mismatch remain measured zero.
An answer range rejected by the official parser also keeps its official-equivalent zero;
this compatibility rule does not establish that any cells were compared.
"""
from __future__ import annotations

import argparse
import datetime
import json
import os
import shutil
import subprocess
import sys
import tempfile
import traceback
from pathlib import Path

DEFAULT_REWARD_FILE = "/verifier/reward.txt"
DEFAULT_EVIDENCE_FILE = "/verifier/evidence.json"


# ------------------------------------------------------- the official scorer ---
# `evaluation/evaluation.py`, transcribed. The function bodies are upstream's; the only
# changes are `print` -> return values, the module-level imports it assumed, and raising
# measurement failures instead of returning the same False as a completed mismatch.


def transform_value(v):
    """`evaluation.py::transform_value`, verbatim."""
    if isinstance(v, (int, float)):
        v = round(float(v), 2)
    elif isinstance(v, datetime.time):
        v = str(v)[:-3]
    elif isinstance(v, datetime.datetime):
        v = round(datetime_to_float(v), 0)
    elif isinstance(v, str):
        try:
            v = round(float(v), 2)
        except ValueError:
            pass
    return v


def datetime_to_float(dt):
    """`evaluation.py::datetime_to_float`, verbatim."""
    excel_start_date = datetime.datetime(1899, 12, 30)
    delta = dt - excel_start_date
    return delta.days + delta.seconds / 86400.0


def compare_cell_value(v1, v2):
    """`evaluation.py::compare_cell_value`, verbatim."""
    v1 = transform_value(v1)
    v2 = transform_value(v2)
    if (v1 == "" and v2 is None) or (v1 is None and v2 == ""):
        return True
    if (v1 == "" and v2 == "") or (v1 is None and v2 is None):
        return True
    if type(v1) != type(v2):
        return False
    if v1 == v2:
        return True
    else:
        return False


def col_num2name(n):
    """`evaluation.py::col_num2name`, verbatim."""
    name = ""
    while n > 0:
        n, remainder = divmod(n - 1, 26)
        name = chr(65 + remainder) + name
    return name


def col_name2num(name):
    """`evaluation.py::col_name2num`, verbatim."""
    num = 0
    for c in name:
        num = num * 26 + (ord(c) - ord("A") + 1)
    return num


def parse_cell_range(range_str):
    """`evaluation.py::parse_cell_range`, verbatim."""
    start_cell, end_cell = range_str.split(":")
    start_col, start_row = "", ""
    for char in start_cell:
        if char.isdigit():
            start_row += char
        else:
            start_col += char
    end_col, end_row = "", ""
    for char in end_cell:
        if char.isdigit():
            end_row += char
        else:
            end_col += char
    return (col_name2num(start_col), int(start_row)), (col_name2num(end_col), int(end_row))


def generate_cell_names(range_str):
    """`evaluation.py::generate_cell_names`, verbatim."""
    if ":" not in range_str:
        return [range_str]
    (start_col, start_row), (end_col, end_row) = parse_cell_range(range_str)
    columns = [col_num2name(i) for i in range(start_col, end_col + 1)]
    return [f"{col}{row}" for col in columns for row in range(start_row, end_row + 1)]


def cell_level_compare(wb_gt, wb_proc, sheet_name, cell_range):
    """`evaluation.py::cell_level_compare`, with the comparisons that are commented out
    upstream left out here too. Returns `(ok, message)`."""
    # Invalid reference metadata is not an agent's missing-sheet mismatch.
    ws_gt = wb_gt[sheet_name]
    if sheet_name not in wb_proc:
        return False, f"worksheet {sheet_name!r} not found in the result"
    ws_proc = wb_proc[sheet_name]
    for cell_name in generate_cell_names(cell_range):
        cell_gt = ws_gt[cell_name]
        cell_proc = ws_proc[cell_name]
        if not compare_cell_value(cell_gt.value, cell_proc.value):
            return False, (f"value difference at {cell_gt.coordinate}: golden has "
                           f"{cell_gt.value!r}, result has {cell_proc.value!r}")
    return True, ""


def compare_workbooks(gt_file, proc_file, instruction_type, answer_position):
    """`evaluation.py::compare_workbooks`, with measurement failures raised.

    `instruction_type` is accepted and unused, exactly as upstream: the comparison does not
    branch on it (upstream only uses it to group the report).
    """
    import openpyxl

    if not os.path.exists(proc_file):
        # main() handles an absent submission; disappearance during comparison is a failure.
        raise FileNotFoundError(f"result workbook unavailable for comparison: {proc_file}")
    try:
        wb_gt = openpyxl.load_workbook(filename=gt_file, data_only=True)
        wb_proc = openpyxl.load_workbook(filename=proc_file, data_only=True)
    except Exception as exc:  # noqa: BLE001 - loading failed, no comparison was completed
        raise RuntimeError(f"workbook load failed: {type(exc).__name__}: {exc}") from exc

    sheet_cell_ranges = answer_position.split(",")
    result_list = []
    msg_list = []
    for sheet_cell_range in sheet_cell_ranges:
        # Compatibility decision (w-spread2, option a): upstream evaluation() catches
        # range-parse ValueError and records int(False). Preserve that measured zero
        # for released metadata such as 283-32; it is not a local runtime failure.
        # Unexpected parser exceptions still propagate to the unmeasured path.
        _cell = sheet_cell_range.split("!")[-1].strip().strip("'")
        # Only a **range** goes through `parse_cell_range`; a single cell (`B6`) is returned
        # as-is by `generate_cell_names` and never parsed. Measured: a first version validated
        # every position, which rejected `B6` and marked `51-12` unmeasurable.
        if ":" in _cell:
            try:
                _ = parse_cell_range(_cell)
            except ValueError as exc:
                return False, (
                    f"official-equivalent measured zero: the task's answer_position "
                    f"{sheet_cell_range!r} is not a cell range the official parser can read: "
                    f"{exc}")
        if "!" in sheet_cell_range:
            sheet_name, cell_range = sheet_cell_range.split("!")
            sheet_name = sheet_name.lstrip("'").rstrip("'")
        else:
            sheet_name = wb_gt.sheetnames[0]
            cell_range = sheet_cell_range
        sheet_name = sheet_name.lstrip("'").rstrip("'")
        cell_range = cell_range.lstrip("'").rstrip("'")
        result, msg = cell_level_compare(wb_gt, wb_proc, sheet_name, cell_range)
        result_list.append(result)
        msg_list.append(msg)
    first_bad = next((m for r, m in zip(result_list, msg_list) if not r), "")
    return all(result_list), first_bad


# ------------------------------------------------------------- recalculation ---


def soffice_path() -> str | None:
    """The LibreOffice binary, from `PATH` or the two conventional names."""
    for candidate in ("soffice", "libreoffice"):
        found = shutil.which(candidate)
        if found:
            return found
    return None


def recalculate(workbook: Path, *, scratch: Path, soffice: str) -> tuple[bool, str]:
    """Re-evaluate `workbook`'s formulas and cache the results, in place.

    LibreOffice's converter writes a file whose cells carry the computed values, which is what
    `data_only=True` then reads. A private `UserInstallation` keeps concurrent checks from
    fighting over one profile lock -- measured: without it, two LibreOffice runs on one host
    can make the second fail.
    """
    profile = scratch / "lo-profile"
    profile.mkdir(parents=True, exist_ok=True)
    outdir = scratch / "conv"
    outdir.mkdir(parents=True, exist_ok=True)
    env = dict(os.environ, HOME=str(scratch),
               TMPDIR=str(scratch))
    cmd = [soffice, "--headless", "--calc", "--norestore", "--nolockcheck",
           f"-env:UserInstallation=file://{profile}",
           "--convert-to", "xlsx:Calc MS Excel 2007 XML",
           "--outdir", str(outdir), str(workbook)]
    try:
        proc = subprocess.run(cmd, capture_output=True, text=True, timeout=180, env=env)
    except subprocess.TimeoutExpired:
        return False, "LibreOffice timed out after 180s"
    if proc.returncode != 0:
        return False, f"LibreOffice exited {proc.returncode}: {proc.stderr.strip()[:300]}"
    converted = outdir / (workbook.stem + ".xlsx")
    if not converted.is_file():
        return False, (f"LibreOffice produced no file (stderr: {proc.stderr.strip()[:200]})")
    shutil.move(str(converted), str(workbook))
    return True, ""


def workbook_has_formulas(path: Path) -> bool:
    """Does any cell hold a formula? Decides whether recalculation is needed."""
    import openpyxl

    try:
        wb = openpyxl.load_workbook(filename=path, data_only=False)
    except Exception as exc:  # noqa: BLE001 - False must mean a readable, formula-free file
        raise RuntimeError(f"workbook load failed during formula detection: "
                           f"{type(exc).__name__}: {exc}") from exc
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    return True
    return False


# -------------------------------------------------------------------- main ---


def parse_args(argv):
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--workdir", default=".")
    parser.add_argument("--answer", default="answer.txt")
    parser.add_argument("--result", default="result.xlsx")
    parser.add_argument("--golden", default="golden.xlsx")
    parser.add_argument("--task", default="task.json")
    parser.add_argument("--reward-file", default=DEFAULT_REWARD_FILE)
    parser.add_argument("--evidence-file", default=DEFAULT_EVIDENCE_FILE)
    args = parser.parse_args(argv)
    args.workdir = Path(args.workdir)
    args.reward_file = Path(args.reward_file)
    args.evidence_file = Path(args.evidence_file)
    return args


def write_reward(path: Path, score: float) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(f"{score:g}\n")


def write_evidence(path: Path, evidence: dict) -> None:
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(evidence, indent=1, sort_keys=True,
                                   ensure_ascii=False) + "\n")
    except OSError as exc:
        print(f"could not write evidence: {exc}", file=sys.stderr)


def main(argv) -> int:
    args = parse_args(argv)
    workdir = Path(args.workdir)
    task_path = workdir / args.task
    golden_path = workdir / args.golden
    answer_path = workdir / args.answer
    evidence: dict = {"answer_file": str(answer_path), "golden": str(golden_path)}

    if not task_path.is_file():
        evidence.update(measured=False, error=f"no task file at {task_path}")
        write_evidence(args.evidence_file, evidence)
        print(f"check cannot measure: {evidence['error']}", file=sys.stderr)
        return 3
    if not golden_path.is_file():
        evidence.update(measured=False, error=f"no golden workbook at {golden_path}")
        write_evidence(args.evidence_file, evidence)
        print(f"check cannot measure: {evidence['error']}", file=sys.stderr)
        return 3

    spec = json.loads(task_path.read_text())
    evidence.update(task_id=spec.get("task_id"),
                    answer_position=spec.get("answer_position"),
                    instruction_type=spec.get("instruction_type"))

    raw_answer = answer_path.read_text(encoding="utf-8").strip() if answer_path.is_file() else ""
    # The answer is the **name of the workbook** the agent left behind. An empty answer means
    # the conventional name; anything else is taken as a relative path inside the workdir, so a
    # harness cannot name a file outside the task.
    named = raw_answer.splitlines()[0].strip() if raw_answer else ""
    named = named or args.result
    candidate = (workdir / named).resolve()
    if not str(candidate).startswith(str(workdir.resolve()) + os.sep):
        evidence.update(measured=False,
                        error=f"the answer names {named!r}, which is outside the workdir")
        write_evidence(args.evidence_file, evidence)
        print(f"check cannot measure: {evidence['error']}", file=sys.stderr)
        return 3
    evidence.update(named_result=named, result_exists=candidate.is_file())

    if not candidate.is_file():
        # The agent left no workbook: a measured zero, not an unmeasured task. It had a
        # contract, and the artifact it names is absent.
        evidence.update(measured=True, reward=0.0,
                        why=f"the answer names {named!r} but no such file exists",
                        score=0.0, correct=False)
        write_evidence(args.evidence_file, evidence)
        write_reward(args.reward_file, 0.0)
        print(f"reward 0: {evidence['why']}")
        return 1

    phase = "formula detection"
    try:
        # Recalculate before comparing: data_only=True reads cached formula values.
        if workbook_has_formulas(candidate):
            phase = "LibreOffice recalculation"
            soffice = soffice_path()
            if soffice is None:
                raise RuntimeError("the result workbook holds formulas and no LibreOffice is "
                                   "installed, so its values cannot be computed; the check "
                                   "will not score an unevaluated workbook")
            with tempfile.TemporaryDirectory(prefix="sb-recalc-") as scratch:
                ok, why = recalculate(candidate, scratch=Path(scratch), soffice=soffice)
            evidence["recalculated"] = ok
            if not ok:
                raise RuntimeError(f"could not recalculate the workbook: {why}")
        else:
            evidence["recalculated"] = False

        phase = "comparison"
        ok, message = compare_workbooks(str(golden_path), str(candidate),
                                        spec.get("instruction_type"),
                                        str(spec["answer_position"]))
    except Exception as exc:  # noqa: BLE001 - a failed tool/comparator is not a zero
        evidence.update(measured=False, error=f"{phase} failed: {type(exc).__name__}: {exc}",
                        traceback=traceback.format_exc()[-2000:])
        write_evidence(args.evidence_file, evidence)
        print(f"check cannot measure: {evidence['error']}", file=sys.stderr)
        return 3

    evidence.update(measured=True, reward=1.0 if ok else 0.0, correct=bool(ok),
                    score=1.0 if ok else 0.0,
                    why=("the named cells hold the reference values" if ok
                         else message if message.startswith("official-equivalent measured zero:")
                         else f"workbook differs: {message or 'comparison failed'}"))
    write_evidence(args.evidence_file, evidence)
    write_reward(args.reward_file, 1.0 if ok else 0.0)
    print(f"reward {1.0 if ok else 0.0:g}: {evidence['why']}")
    return 0 if ok else 1


if __name__ == "__main__":
    try:
        raise SystemExit(main(sys.argv[1:]))
    except SystemExit:
        raise
    except Exception:  # noqa: BLE001 - a crash must not look like a measured zero
        traceback.print_exc()
        raise SystemExit(4)

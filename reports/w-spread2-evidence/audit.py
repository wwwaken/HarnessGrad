"""Metadata/parser audit plus synthetic upstream caller reproduction; no benchmark run."""
from __future__ import annotations

import ast
import base64
from contextlib import contextmanager
import datetime
import hashlib
import io
import json
import os
from pathlib import Path
import tarfile
from types import SimpleNamespace

import openpyxl

HERE = Path(__file__).resolve().parent
SOURCE = HERE / 'official-evaluation.py.txt'
DATA = Path('/mnt/20t/xubuqiang/datasets/spreadsheetbench')
METADATA = DATA / 'spreadsheetbench_verified_400/dataset.json'
ARCHIVE = DATA / '_src/spreadsheetbench_verified_400.tar.gz'
REVISION = 'e2758ac01c6c52efc746bf0d53b2fd01dbc42dd2'


def official_functions():
    # No upstream imports, CLI, runtime, or main block is executed.
    tree = ast.parse(SOURCE.read_text(), filename=str(SOURCE))
    selected = {'datetime_to_float', 'transform_value', 'compare_cell_value',
                'col_num2name', 'col_name2num', 'parse_cell_range', 'generate_cell_names',
                'cell_level_compare', 'compare_workbooks', 'evaluation'}
    module = ast.Module(body=[node for node in tree.body
                             if isinstance(node, ast.FunctionDef) and node.name in selected],
                        type_ignores=[])
    namespace = {'datetime': datetime, 'json': json, 'os': os, 'openpyxl': openpyxl,
                 'tqdm': lambda values: values}
    exec(compile(module, str(SOURCE), 'exec'), namespace)
    return namespace


def main():
    response = json.loads((HERE / 'official-source-response.json').read_text())
    source = base64.b64decode(response['content'])
    # Keep raw upstream bytes losslessly in the API response. The readable source
    # view trims trailing whitespace only; its AST and source line numbers agree.
    view = '\n'.join(line.rstrip() for line in source.decode().splitlines()).rstrip() + '\n'
    assert SOURCE.read_text() == view
    assert ast.dump(ast.parse(source)) == ast.dump(ast.parse(view))
    blob = hashlib.sha1(b'blob ' + str(len(source)).encode() + b'\0' + source).hexdigest()
    assert blob == response['sha']
    print('upstream_revision:', REVISION)
    print('source_url:', response['html_url'])
    print('source_sha256:', hashlib.sha256(source).hexdigest())
    print('source_git_blob_sha1:', blob)
    metadata = METADATA.read_bytes()
    archive_hash = hashlib.sha256(ARCHIVE.read_bytes()).hexdigest()
    with tarfile.open(ARCHIVE) as archive:
        member = archive.extractfile('spreadsheetbench_verified_400/dataset.json').read()
    print('archive_sha256:', archive_hash)
    print('matches_adapter_declared_archive_sha256:', archive_hash ==
          '10ef893dd29cb13ab97143ea787e68cdc9574a13873ab9a54e50b31dc03fc949')
    print('metadata_path:', str(METADATA))
    print('metadata_sha256:', hashlib.sha256(metadata).hexdigest())
    print('metadata_identical_to_local_archive_member:', metadata == member)
    rows = json.loads(metadata)
    row = next(row for row in rows if str(row['id']) == '283-32')
    print('raw_metadata_fields:', json.dumps({key: row[key] for key in
          ('id', 'spreadsheet_path', 'instruction_type', 'answer_position')}))
    for number, line in enumerate(metadata.decode().splitlines(), 1):
        if line.strip() in ('"id": "283-32",',
                            '"answer_position": "Sheet3\'!A:G,\'Sheet4\'!A:G",'):
            print(f'raw_metadata_line_{number}: {line}')
    namespace = official_functions()
    bad = []
    colon_ranges = single_cells = 0
    for record in rows:
        for position in str(record['answer_position']).split(','):
            # Exactly the official quote handling; don't strip additional whitespace.
            cell_range = position.split('!')[-1].lstrip("'").rstrip("'")
            if ':' not in cell_range:
                single_cells += 1
                continue
            colon_ranges += 1
            try:
                namespace['parse_cell_range'](cell_range)
            except Exception as exc:
                bad.append({'id': str(record['id']), 'range': cell_range,
                            'exception': f'{type(exc).__name__}: {exc}'})
    print('audit_method: every comma-separated component; official parse_cell_range for colon ranges; single cells bypass it upstream')
    print('metadata_rows:', len(rows))
    print('colon_range_components:', colon_ranges)
    print('single_cell_components:', single_cells)
    print('parse_failing_components:', json.dumps(bad))
    print('parse_failing_tasks:', len({entry['id'] for entry in bad}))
    print('scope: installed verified_400 metadata only; not a gradability or full upstream 912-task audit')
    print('283-32_hash_first_byte:', hashlib.sha256(b'283-32').digest()[0])
    print('283-32_default_hash_eval:', hashlib.sha256(b'283-32').digest()[0] % 3 == 0)

    # Run the official functions on in-memory authored workbooks only. The real task's
    # workbook/instruction is never executed or scored. Reuse only its malformed range.
    workbook = openpyxl.Workbook()
    workbook.active.title = 'Sheet3'
    workbook.create_sheet('Sheet4')
    namespace['openpyxl'] = SimpleNamespace(load_workbook=lambda **kwargs: workbook)
    namespace['os'] = SimpleNamespace(path=SimpleNamespace(
        exists=lambda path: True, abspath=lambda path: '/synthetic-only'))
    position = row['answer_position']
    try:
        namespace['compare_workbooks']('synthetic-reference', 'synthetic-result', '', position)
    except Exception as exc:
        print('official_compare_on_synthetic_workbooks:', f'{type(exc).__name__}: {exc}')
    else:
        raise AssertionError('official comparator should raise on this range')
    output = io.StringIO()
    synthetic_rows = [{'id': 'synthetic-released-range', 'instruction_type': '',
                       'answer_position': position}]
    @contextmanager
    def synthetic_open(path, mode):
        if mode == 'r':
            yield io.StringIO(json.dumps(synthetic_rows))
        else:
            assert mode == 'w'
            yield output
    namespace['open'] = synthetic_open
    namespace['evaluation'](SimpleNamespace(dataset='synthetic-only', setting='single',
                                          model='no-model-executed'))
    actual = json.loads(output.getvalue())[0]
    assert actual['test_case_results'] == [0, 0, 0]
    assert actual['soft_restriction'] == actual['hard_restriction'] == 0
    print('official_evaluation_on_synthetic_workbooks:', output.getvalue())
    workbook.close()


if __name__ == '__main__':
    main()

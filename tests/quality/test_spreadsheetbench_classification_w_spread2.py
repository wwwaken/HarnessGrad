"""Synthetic checker/admission regressions; no dataset payload or external runtime."""
from __future__ import annotations

from contextlib import contextmanager, redirect_stderr, redirect_stdout
import importlib.util
import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
from types import SimpleNamespace
import zipfile
from xml.etree import ElementTree as ET

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from eval.run import runner
from eval.verification import verifier


def _module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


CASES = [
    ("unparseable-range", 0.0, "not a cell range"),
    ("released-unparseable-range", 0.0, "official-equivalent measured zero"),
    ("parser-exception", None, "synthetic parser failure"),
    ("result-vanished", None, "unavailable for comparison"),
    ("golden-load", None, "workbook load failed"),
    ("result-load", None, "workbook load failed"),
    ("formula-probe-load", None, "formula detection"),
    ("comparator-exception", None, "synthetic comparator failure"),
    ("formula-no-runtime", None, "no LibreOffice"),
    ("recalc-timeout", None, "LibreOffice timed out"),
    ("recalc-exit", None, "LibreOffice exited 7"),
    ("recalc-no-output", None, "LibreOffice produced no file"),
    ("recalc-unevaluated-copy", None, "did not populate a fresh formula cache"),
    ("recalc-partial-cache", None, "did not populate a fresh formula cache"),
    ("recalc-array-cache", None, "cannot observe fresh caches for non-scalar formulas"),
    ("recalc-launch", None, "LibreOffice recalculation"),
    ("recalc-copy", None, "LibreOffice recalculation"),
    ("recalc-scratch", None, "LibreOffice recalculation"),
    ("recalc-cleanup", None, "LibreOffice recalculation"),
    ("golden-sheet-absent", None, "comparison"),
    ("mismatch", 0.0, "value difference"),
    ("result-sheet-absent", 0.0, "not found in the result"),
    ("result-absent", 0.0, "no such file exists"),
    ("match", 1.0, "reference values"),
]


@pytest.fixture(params=CASES, ids=[c[0] for c in CASES])
def scenario(request, monkeypatch, tmp_path):
    case, score, reason = request.param
    check = _module("spreadsheet_check_w_spread2",
                    ROOT / "data/datasets/spreadsheetbench/verify/check.py")
    # Import the adapter against an empty synthetic root, never the installed mirror.
    monkeypatch.setenv("HG_SPREADSHEETBENCH_DATA", str(tmp_path / "empty-data"))
    adapter = _module("spreadsheet_adapter_w_spread2", ROOT / "data/datasets/spreadsheetbench.py")
    work = tmp_path / "work"
    work.mkdir()
    golden = tmp_path / "reference.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet3" if case == "released-unparseable-range" else "S"
    wb.active["A1"] = 1
    wb.save(golden)
    if case == "mismatch":
        wb.active["A1"] = 2
    elif case.startswith("recalc-") or case == "formula-no-runtime":
        wb.active["A1"] = "=1"
        if case == "recalc-partial-cache":
            wb.active["A2"] = "=2"
    if case == "result-sheet-absent":
        wb.active.title = "Other"
    if case != "result-absent":
        wb.save(work / "result.xlsx")
    wb.close()
    if case == "recalc-array-cache":
        candidate = work / "result.xlsx"
        with zipfile.ZipFile(candidate) as src:
            entries = {name: src.read(name) for name in src.namelist()}
        entries["xl/worksheets/sheet1.xml"] = entries[
            "xl/worksheets/sheet1.xml"].replace(b"<f>", b'<f t="array" ref="A1:A2">')
        with zipfile.ZipFile(candidate, "w") as dst:
            for name, content in entries.items():
                dst.writestr(name, content)
    if case == "golden-load":
        golden.write_bytes(b"synthetic unreadable workbook")
    if case == "result-load":
        (work / "result.xlsx").write_bytes(b"synthetic unreadable workbook")
    position = ("Sheet3'!A:G,'Sheet4'!A:G" if case == "released-unparseable-range" else
                "S!A:G" if case == "unparseable-range" else
                "Absent!A1" if case == "golden-sheet-absent" else "S!A1")
    record = {"task_id": "283-32" if case == "released-unparseable-range" else "synthetic-spread", "golden_path": golden,
              "instruction_type": "", "answer_position": position}
    spec = adapter._verifier(record)
    (work / "task.json").write_text(json.dumps(record, default=str))
    (work / "golden.xlsx").write_bytes(golden.read_bytes())
    if case == "formula-probe-load":
        load = openpyxl.load_workbook

        def probe_load(*args, **kwargs):
            if kwargs.get("data_only") is False:
                raise OSError("synthetic formula probe load failure")
            return load(*args, **kwargs)

        monkeypatch.setattr(openpyxl, "load_workbook", probe_load)
    if case == "parser-exception":
        def broken_parser(*args, **kwargs):
            raise RuntimeError("synthetic parser failure")
        monkeypatch.setattr(check, "parse_cell_range", broken_parser)
        record["answer_position"] = "S!A1:A1"
        (work / "task.json").write_text(json.dumps(record, default=str))
        spec = adapter._verifier(record)
    if case == "result-vanished":
        compare = check.compare_workbooks
        def vanished_result(gt_file, proc_file, *args):
            Path(proc_file).unlink()
            return compare(gt_file, proc_file, *args)
        monkeypatch.setattr(check, "compare_workbooks", vanished_result)
    if case in {"recalc-scratch", "recalc-cleanup"}:
        temporary_directory = tempfile.TemporaryDirectory
        @contextmanager
        def failed_scratch(*args, **kwargs):
            if case == "recalc-scratch":
                raise OSError("synthetic scratch setup failure")
            with temporary_directory(*args, **kwargs) as scratch:
                yield scratch
            raise OSError("synthetic scratch cleanup failure")
        # Replace the *check's* reference to the module, not the attribute on the shared
        # `tempfile` module. `check.tempfile` **is** `tempfile`, so patching the attribute
        # also broke `data/scaffolds/scorer_input.py:77`, whose own TemporaryDirectory is
        # what validates the scorer-input receipt before the check scores anything. The
        # check then crashed with exit 4 before it could write evidence, and the reason the
        # test was looking for never reached the verdict. This injects the fault at the
        # check's scratch step and nowhere else.
        monkeypatch.setattr(check, "tempfile",
                            SimpleNamespace(TemporaryDirectory=failed_scratch))
    if case == "comparator-exception":
        def broken_compare(*args, **kwargs):
            raise RuntimeError("synthetic comparator failure")
        monkeypatch.setattr(check, "compare_workbooks", broken_compare)
    monkeypatch.setattr(check, "soffice_path",
                        lambda: None if case == "formula-no-runtime" else "synthetic-soffice")

    def denied(*args, **kwargs):
        pytest.fail("unexpected external runtime launch")

    monkeypatch.setattr(verifier.sandbox_mod, "run", denied)
    monkeypatch.setattr(verifier.container_mod, "run", denied)
    monkeypatch.setattr(verifier.container_mod, "create", denied)

    def invoke(argv):
        stdout, stderr = io.StringIO(), io.StringIO()
        with redirect_stdout(stdout), redirect_stderr(stderr):
            try:
                code = check.main(argv)
            except Exception:
                # Match check.py's __main__ crash exit, so pre-fix missing evidence is tested.
                code = 4
        return subprocess.CompletedProcess(["synthetic-check"], code,
                                           stdout.getvalue(), stderr.getvalue())

    def synthetic_process(cmd, **kwargs):
        if cmd[0] == "synthetic-soffice":
            if case == "recalc-timeout":
                raise subprocess.TimeoutExpired(cmd, 180)
            if case == "recalc-launch":
                raise OSError("synthetic LibreOffice launch failure")
            if case in {"recalc-unevaluated-copy", "recalc-partial-cache"}:
                outdir = Path(cmd[cmd.index("--outdir") + 1])
                source = Path(cmd[-1])
                if case == "recalc-unevaluated-copy":
                    shutil.copyfile(source, outdir / "result.xlsx")
                else:
                    with zipfile.ZipFile(source) as src, zipfile.ZipFile(
                            outdir / "result.xlsx", "w") as dst:
                        for entry in src.infolist():
                            content = src.read(entry)
                            if entry.filename == "xl/worksheets/sheet1.xml":
                                sheet = ET.fromstring(content)
                                ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
                                first = sheet.find(f".//{ns}c[@r='A1']")
                                value = first.find(ns + "v")
                                if value is None:
                                    value = ET.SubElement(first, ns + "v")
                                value.text = "1"
                                content = ET.tostring(sheet)
                            dst.writestr(entry, content)
            if case == "recalc-copy":
                # Populate a valid cache first, then fail only the final publication move.
                outdir = Path(cmd[cmd.index("--outdir") + 1])
                workbook = openpyxl.Workbook()
                workbook.active["A1"] = "=1"
                workbook.save(outdir / "result.xlsx")
                workbook.close()
                with zipfile.ZipFile(outdir / "result.xlsx") as src:
                    entries = {name: src.read(name) for name in src.namelist()}
                entries["xl/worksheets/sheet1.xml"] = entries[
                    "xl/worksheets/sheet1.xml"].replace(b"<v></v>", b"<v>1</v>")
                with zipfile.ZipFile(outdir / "result.xlsx", "w") as dst:
                    for name, content in entries.items():
                        dst.writestr(name, content)
                monkeypatch.setattr(check.shutil, "move", denied_move)
            return subprocess.CompletedProcess(cmd, 7 if case == "recalc-exit" else 0,
                                               "", "synthetic recalculation stderr")
        assert cmd[:2] == ["python3", "check.py"], cmd
        assert (Path(kwargs["cwd"]) / "check.py").read_bytes() == (
            ROOT / "data/datasets/spreadsheetbench/verify/check.py").read_bytes()
        argv = list(cmd[2:])
        argv[argv.index("--workdir") + 1] = str(kwargs["cwd"])
        for flag in ("--reward-file", "--evidence-file"):
            i = argv.index(flag) + 1
            argv[i] = str(Path(kwargs["cwd"]) / argv[i])
        return invoke(argv)

    def denied_move(*args, **kwargs):
        raise OSError("synthetic conversion copy failure")

    monkeypatch.setattr(verifier.subprocess, "run", synthetic_process)
    return check, work, spec, score, reason, invoke, record["task_id"]


def test_checker_classifies_failures_and_completed_comparisons(scenario):
    check, work, spec, score, reason, invoke, task_id = scenario
    proc = invoke(["--workdir", str(work), "--reward-file", str(work / "reward.txt"),
                   "--evidence-file", str(work / "evidence.json")])
    evidence = json.loads((work / "evidence.json").read_text())
    if score is None:
        assert proc.returncode == 3
        assert evidence["measured"] is False
        if reason.startswith(("did not populate", "cannot observe fresh")):
            assert evidence["recalculated"] is False
        assert reason in evidence["error"]
        assert not {"score", "reward", "correct"}.intersection(evidence)
        assert not (work / "reward.txt").exists()
        assert "reward 0" not in proc.stdout and "reward 1" not in proc.stdout
    else:
        assert proc.returncode == (0 if score else 1)
        assert evidence["measured"] is True
        assert evidence["score"] == evidence["reward"] == score
        assert reason in evidence["why"]
        assert float((work / "reward.txt").read_text()) == score


def test_classification_survives_verifier_persistence_and_aggregation(
        scenario, monkeypatch, tmp_path):
    check, work, spec, score, reason, invoke, task_id = scenario
    verdict = verifier._verify(work, spec, "", sandbox=False, timeout_s=300,
                               env={"kind": "files"}, task_id=task_id)
    monkeypatch.setattr(runner, "run_one", lambda *a, **k: {
        "verdict": verdict, "trace": "synthetic trace\n", "exit_code": 0, "stderr": ""})
    cache = {}
    result = runner._evaluate(
        tmp_path, [{"task_id": task_id}], {}, cache=cache,
        trials=1, jobs=1, progress_dir=tmp_path, round_no=0)
    persisted = json.loads((tmp_path / f"progress/round-0/{task_id}.json").read_text())
    if score is None:
        assert reason in verdict["unmeasured"]
        assert result["per_task"] == result["per_task_trials"] == {}
        assert result["trial_verdicts"] == result["verdicts"] == {}
        assert result["invalid"][task_id]["stage"] == "check"
        assert reason in persisted["invalid"]["detail"]
        assert not {"score", "trial_scores", "complete"}.intersection(persisted)
        assert cache == {}
    else:
        assert "unmeasured" not in verdict
        assert result["per_task"] == {task_id: score}
        assert result["invalid"] == {}
        assert persisted["complete"] is True and persisted["score"] == score
        assert len(cache) == 1


@pytest.mark.skipif(os.environ.get("HG_SPREADCACHE_ENGINE_CONTROL") != "1",
                    reason="opt-in authored real LibreOffice control (no dataset/model)")
def test_real_engine_replaces_authored_stale_caches(tmp_path):
    """The original failing mathematical assertion, under explicit network isolation."""
    assert shutil.which("bwrap") and shutil.which("soffice")
    wb = openpyxl.Workbook()
    wb.active.title = "Synthetic"
    wb.active["A1"], wb.active["A2"] = 2, 3
    wb.active["B1"], wb.active["B2"] = "=SUM(A1:A2)", "=B1*2"
    candidate = tmp_path / "result.xlsx"
    wb.save(candidate)
    with zipfile.ZipFile(candidate) as src:
        entries = {n: src.read(n) for n in src.namelist()}
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    sheet = ET.fromstring(entries["xl/worksheets/sheet1.xml"])
    for name, stale in (("B1", 999), ("B2", 888)):
        sheet.find(f".//{ns}c[@r='{name}']/{ns}v").text = str(stale)
    entries["xl/worksheets/sheet1.xml"] = ET.tostring(sheet)
    with zipfile.ZipFile(candidate, "w", zipfile.ZIP_DEFLATED) as dst:
        for name, content in entries.items():
            dst.writestr(name, content)
    wb.active["B1"], wb.active["B2"] = 5, 10
    wb.save(tmp_path / "golden.xlsx")
    wb.close()
    cached = openpyxl.load_workbook(candidate, data_only=True)
    assert [cached.active["B1"].value, cached.active["B2"].value] == [999, 888]
    cached.close()
    (tmp_path / "task.json").write_text(json.dumps({
        "task_id": "spreadcache-authored", "answer_position": "Synthetic!B1:B2"}))
    (tmp_path / "answer.txt").write_text("result.xlsx\n")
    scratch = tmp_path / "scratch"
    scratch.mkdir()
    cmd = [shutil.which("bwrap"), "--unshare-net", "--ro-bind", "/", "/",
           "--bind", str(tmp_path), "/tmp",
           "--proc", "/proc", "--", sys.executable,
           str(ROOT / "data/datasets/spreadsheetbench/verify/check.py"),
           "--workdir", "/tmp", "--reward-file", "/tmp/reward.txt",
           "--evidence-file", "/tmp/evidence.json"]
    proc = subprocess.run(cmd, capture_output=True, text=True, timeout=240,
                          env=dict(os.environ, TMPDIR="/tmp/scratch", PYTHONDONTWRITEBYTECODE="1"))
    assert proc.returncode == 0, (proc.stdout, proc.stderr)
    cached = openpyxl.load_workbook(candidate, data_only=True)
    assert [cached.active["B1"].value, cached.active["B2"].value] == [5, 10]
    cached.close()
    evidence = json.loads((tmp_path / "evidence.json").read_text())
    assert evidence["recalculated"] is True and evidence["measured"] is True
    assert evidence["score"] == 1.0 and (tmp_path / "reward.txt").read_text() == "1\n"

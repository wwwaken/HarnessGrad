#!/usr/bin/env python3
"""Read-only structural inventory. Never imports or invokes a scorer/comparator.

Only allowlisted parser/naming function ASTs are compiled from shipped source.
Workbook values and formulas are inspected for types, never compared or exported.
"""
import argparse
import ast
from collections import Counter
from io import BytesIO
import hashlib
import json
from pathlib import Path
import posixpath
import sys
import tarfile
import warnings
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import openpyxl


ROOT = Path(__file__).resolve().parents[3]
DATA = Path('/mnt/20t/xubuqiang/datasets/spreadsheetbench')
RELEASE = 'spreadsheetbench_verified_400'
META = f'{RELEASE}/dataset.json'
ARCHIVE = DATA / '_src/spreadsheetbench_verified_400.tar.gz'
EXPECTED_META = 'bcecaa89a005bd4e3bbe98da150a86e8062c27f262e575d5e47bd9861b3525e7'
EXPECTED_ARCHIVE = '10ef893dd29cb13ab97143ea787e68cdc9574a13873ab9a54e50b31dc03fc949'
CHECK = 'data/datasets/spreadsheetbench/verify/check.py'
ADAPTER = 'data/datasets/spreadsheetbench.py'
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def sha(b):
    return hashlib.sha256(b).hexdigest()


def canonical(v):
    return json.dumps(v, ensure_ascii=False, sort_keys=True,
                      separators=(',', ':'), allow_nan=False).encode()


def extract(path, names, namespace):
    tree = ast.parse((ROOT / path).read_bytes())
    selected = [n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name in names]
    assert {n.name for n in selected} == set(names)
    exec(compile(ast.Module(body=selected, type_ignores=[]), path, 'exec'), namespace)
    return {n.name: f'{path}:{n.lineno}-{n.end_lineno}' for n in selected}


def read_book(path, data, components):
    """Structural probes over one already-hashed immutable byte buffer."""
    out = {'path': str(path.relative_to(DATA)), 'sha256': sha(data), 'ooxml_open': False,
           'xml_readable': False, 'cached_open': False, 'formula_open': False}
    try:
        with ZipFile(BytesIO(data)) as z:
            # Read every ZIP member (CRC checked); parse every XML/.rels member.
            xml_formula_types = Counter()
            for name in z.namelist():
                b = z.read(name)
                if name.endswith(('.xml', '.rels')):
                    node = ET.fromstring(b)
                    if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                        for f in node.findall('.//s:f', NS):
                            xml_formula_types[f.get('t', 'normal')] += 1
            rels = {r.get('Id'): r.get('Target')
                    for r in ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))}
            physical_cells = {}
            for sheet in ET.fromstring(z.read('xl/workbook.xml')).findall('s:sheets/s:sheet', NS):
                rid = sheet.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                target = rels[rid]
                member = (target.lstrip('/') if target.startswith('/') else
                          posixpath.normpath(posixpath.join('xl', target)))
                physical_cells[sheet.get('name')] = {
                    c.get('r') for c in ET.fromstring(z.read(member)).findall('.//s:c', NS)}
        out['xml_readable'] = True
        out['xml_formula_types'] = dict(sorted(xml_formula_types.items()))
        loaded = []
        load_errors = {}
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter('always', UserWarning)
            warnings.simplefilter('ignore', DeprecationWarning)
            for cached in (False, True):
                key = 'cached_open' if cached else 'formula_open'
                try:
                    loaded.append(openpyxl.load_workbook(BytesIO(data), data_only=cached))
                    out[key] = True
                except Exception as exc:
                    chain = []
                    while exc is not None:
                        chain.append(f'{type(exc).__name__}: {exc}')
                        exc = exc.__cause__
                    load_errors[key] = chain
        out['warnings'] = sorted({str(w.message) for w in caught})
        if load_errors:
            out['load_errors'] = load_errors
            out['error'] = '; '.join(f'{k}: {v[-1]}' for k, v in load_errors.items())
            for book in loaded:
                book.close()
            return out
        formula_book, cache_book = loaded
        out.update(ooxml_open=True, sheets=formula_book.sheetnames,
                   string_formula_cells=0, nonstring_formula_cells=0,
                   targets=[])
        for ws in formula_book.worksheets:
            # _cells is the materialized cell set; implicit blanks cannot start '='.
            for cell in ws._cells.values():
                if isinstance(cell.value, str) and cell.value.startswith('='):
                    out['string_formula_cells'] += 1
                elif cell.data_type == 'f':
                    out['nonstring_formula_cells'] += 1
        for comp in components:
            # Reproduce the checker's first-golden-sheet default in both books.
            sheet = comp.get('sheet') or components[0]['default_golden_sheet']
            item = {'sheet': sheet, 'sheet_present': sheet in formula_book,
                    'addressable': None, 'explicit_ooxml_cells': None,
                    'implicit_cells': None, 'merged_placeholders': None}
            names = comp.get('_cells')
            if names is not None and sheet in formula_book:
                ws = formula_book[sheet]
                serialized = physical_cells[sheet]
                item['explicit_ooxml_cells'] = sum(n in serialized for n in names)
                item['implicit_cells'] = sum(n not in serialized for n in names)
                item['merged_placeholders'] = sum(
                    isinstance(ws._cells.get(openpyxl.utils.cell.coordinate_to_tuple(n)),
                               openpyxl.cell.cell.MergedCell)
                    for n in names if openpyxl.utils.cell.COORD_RE.fullmatch(n))
                # Addressability is not nonempty content or comparison success.
                try:
                    for name in names:
                        ws[name]
                        cache_book[sheet][name]
                    item['addressable'] = bool(names)
                except Exception as exc:
                    item['addressable'] = False
                    item['error'] = f'{type(exc).__name__}: {exc}'
            elif names is not None:
                item['addressable'] = False
            out['targets'].append(item)
        out['formula_route_if_submitted_unchanged'] = (
            'engine_required_array_also_present' if out['string_formula_cells'] and out['nonstring_formula_cells'] else
            'engine_required' if out['string_formula_cells'] else
            'array_only_detector_bypass' if out['nonstring_formula_cells'] else 'no_formula_observed')
        for book in loaded:
            book.close()
    except Exception as exc:
        out['error'] = f'{type(exc).__name__}: {exc}'
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--out', type=Path, required=True)
    args = ap.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)
    parser_ns = {}
    parser_sources = extract(CHECK, ['col_num2name', 'col_name2num', 'parse_cell_range',
                                     'generate_cell_names'], parser_ns)
    loader_ns = {'Path': Path}
    adapter_ast = ast.parse((ROOT / ADAPTER).read_bytes())
    patterns = next(n.value for n in adapter_ast.body if isinstance(n, ast.Assign)
                    and any(isinstance(t, ast.Name) and t.id == '_WORKBOOK_PATTERNS'
                            for t in n.targets))
    loader_ns['_WORKBOOK_PATTERNS'] = ast.literal_eval(patterns)
    loader_sources = extract(ADAPTER, ['_find_workbook'], loader_ns)
    meta_bytes = (DATA / META).read_bytes()
    assert sha(meta_bytes) == EXPECTED_META
    rows = json.loads(meta_bytes)
    assert len(rows) == len({str(r['id']) for r in rows}) == 400
    archive_bytes = ARCHIVE.read_bytes()
    assert sha(archive_bytes) == EXPECTED_ARCHIVE
    records = []
    with tarfile.open(fileobj=BytesIO(archive_bytes), mode='r:gz') as archive:
        assert archive.extractfile(META).read() == meta_bytes
        for index, row in enumerate(rows):
            task = str(row['id'])
            folder = DATA / RELEASE / str(row['spreadsheet_path'])
            paths = {k: loader_ns['_find_workbook'](folder, task, k)
                     for k in ('init', 'golden')}
            components = []
            for raw in str(row.get('answer_position') or '').strip().split(','):
                probe = raw.split('!')[-1].strip().strip("'")
                c = {'component': raw, 'colon': ':' in probe}
                try:
                    if ':' in probe:
                        c['bounds'] = parser_ns['parse_cell_range'](probe)
                        c['parse'] = 'accepted'
                    else:
                        c['parse'] = 'single_cell_bypass'
                except Exception as exc:
                    c.update(parse=type(exc).__name__, error=str(exc))
                if '!' in raw:
                    pieces = raw.split('!')
                    if len(pieces) != 2:
                        c['reference_error'] = 'multiple ! separators'
                        c['_cells'] = None
                        components.append(c)
                        continue
                    sheet, cell_range = pieces
                    c['sheet'] = sheet.lstrip("'").rstrip("'")
                else:
                    cell_range = raw
                c['range'] = cell_range.lstrip("'").rstrip("'")
                try:
                    c['_cells'] = parser_ns['generate_cell_names'](c['range'])
                    c['cell_count'] = len(c['_cells'])
                except Exception as exc:
                    c['_cells'] = None
                    c['reference_error'] = f'{type(exc).__name__}: {exc}'
                components.append(c)
            # Read first sheet name only, without comparator invocation.
            if paths['golden'] is not None:
                b = paths['golden'].read_bytes()
                with warnings.catch_warnings():
                    warnings.simplefilter('ignore')
                    wb = openpyxl.load_workbook(BytesIO(b), read_only=True)
                components[0]['default_golden_sheet'] = wb.sheetnames[0]
                wb.close()
            books = {}
            for kind, path in paths.items():
                if path is None:
                    books[kind] = {'path': None, 'ooxml_open': False, 'error': 'missing workbook'}
                    continue
                data = path.read_bytes()
                member = str(path.relative_to(DATA))
                assert archive.extractfile(member).read() == data, member
                books[kind] = read_book(path, data, components)
                books[kind]['archive_member_equal'] = True
            xlsx = sorted(p.name for p in folder.glob('*.xlsx'))
            is_single = len(xlsx) == 2 and {p.name for p in paths.values() if p} == set(xlsx)
            errors = []
            if not is_single:
                errors.append('workbook inventory is not exactly the loader-selected pair')
            for k, book in books.items():
                if not book['ooxml_open'] or 'error' in book:
                    errors.append(f'{k}: {book.get("error")}')
            policy = (task == '283-32' and len(components) == 2
                      and all(c['range'] == 'A:G' and c['parse'] == 'ValueError' for c in components)
                      and row['answer_position'] == "Sheet3'!A:G,'Sheet4'!A:G")
            if not policy:
                for c in components:
                    if 'reference_error' in c or c.get('cell_count') == 0:
                        errors.append(f'reference: {c.get("reference_error", "empty range")}')
                # Init may intentionally lack a result sheet, but golden must name it.
                for target in books['golden'].get('targets', []):
                    if not target['sheet_present'] or not target['addressable']:
                        errors.append(f'golden target unavailable: {target["sheet"]}: '
                                      f'{target.get("error", "missing sheet or empty expansion")}')
            for c in components:
                c.pop('_cells', None)
            rec = {'task_id': task, 'metadata_index': index,
                   'metadata_sha256': sha(meta_bytes), 'row_sha256': sha(canonical(row)),
                   'default_side': 'eval' if hashlib.sha256(task.encode()).digest()[0] % 3 == 0 else 'train',
                   'answer_position': row['answer_position'], 'answer_sheet': row.get('answer_sheet'),
                   'instruction_type': row.get('instruction_type'),
                   'declared_exclude': row.get('exclude'),
                   'instruction_mentions_array': 'array' in row['instruction'].lower(),
                   'single_variant': is_single, 'xlsx_members': xlsx,
                   'references': components, 'inputs': books,
                   'classification': ('not_gradable_as_installed' if errors else
                                      'range_zero_policy_only' if policy else 'structurally_gradable'),
                   'reason': '; '.join(errors) if errors else
                             ('audited 283-32 A:G ValueError; policy zero only if reached' if policy else
                              'single pair readable; golden target cells addressable; parser accepts or bypasses'),
                   'execution_eligibility': 'UNKNOWN', 'measured': None}
            records.append(rec)
    inventory = b''.join(canonical(r) + b'\n' for r in records)
    (args.out / 'eligibility.jsonl').write_bytes(inventory)
    counts = Counter(r['classification'] for r in records)
    summary = {
        'dataset': 'spreadsheetbench', 'scope': 'installed Verified-400; structural only',
        'metadata': {'path': str(DATA / META), 'sha256': sha(meta_bytes), 'bytes': len(meta_bytes)},
        'archive': {'path': str(ARCHIVE), 'sha256': sha(archive_bytes), 'bytes': len(archive_bytes)},
        'reused_release_authentication': 'reports/w-jc2-review.md:215-239',
        'rows': len(records), 'classes': {k: counts[k] for k in
            ('structurally_gradable', 'range_zero_policy_only', 'not_gradable_as_installed')},
        'default_sides': dict(Counter(r['default_side'] for r in records)),
        'parser_sources': parser_sources, 'loader_sources': loader_sources,
        'source_sha256': {p: sha((ROOT / p).read_bytes()) for p in (CHECK, ADAPTER)},
        'python': {'executable': sys.executable, 'version': sys.version, 'openpyxl': openpyxl.__version__},
        'workbooks_opened': sum(b['ooxml_open'] for r in records for b in r['inputs'].values()),
        'single_variant_rows': sum(r['single_variant'] for r in records),
        'all_workbooks_xml_readable': all(b['xml_readable'] for r in records for b in r['inputs'].values()),
        'declared_exclusions': {r['task_id']: r['declared_exclude'] for r in records if r['declared_exclude']},
        'array_instruction_rows': [r['task_id'] for r in records if r['instruction_mentions_array']],
        'formula_rows': {k: [r['task_id'] for r in records if r['inputs'][k].get('string_formula_cells')]
                         for k in ('init', 'golden')},
        'nonstring_formula_rows': {k: [r['task_id'] for r in records if r['inputs'][k].get('nonstring_formula_cells')]
                                  for k in ('init', 'golden')},
        'xml_array_rows': {k: [r['task_id'] for r in records if r['inputs'][k].get('xml_formula_types', {}).get('array')]
                           for k in ('init', 'golden')},
        'formula_routes': {k: dict(Counter(r['inputs'][k].get('formula_route_if_submitted_unchanged',
                                                            'workbook_load_refusal') for r in records))
                           for k in ('init', 'golden')},
        'eligibility_by_default_side': {side: dict(Counter(r['classification'] for r in records
                                                          if r['default_side'] == side))
                                        for side in ('train', 'eval')},
        'structural_denominator': counts['structurally_gradable'],
        'policy_inclusive_candidate_denominator': counts['structurally_gradable'] + counts['range_zero_policy_only'],
        'denominator_limits': 'Candidate inventory only; future authorized manifest must expose 400 released / '
                             'structural / range-zero / not-gradable counts, resolve exclusions and runtime, '
                             'and bind chosen split and side. No measured or execution-eligible denominator.',
        'reference_components': dict(Counter(c['parse'] for r in records for c in r['references'])),
        'not_gradable_ids': [r['task_id'] for r in records if r['classification'] == 'not_gradable_as_installed'],
        'range_zero_ids': [r['task_id'] for r in records if r['classification'] == 'range_zero_policy_only'],
        'missing_target_sheets': {k: {r['task_id']: [t['sheet'] for t in r['inputs'][k].get('targets', [])
                                                  if not t['sheet_present']]
                                     for r in records if any(not t['sheet_present']
                                         for t in r['inputs'][k].get('targets', []))}
                                 for k in ('init', 'golden')},
        'inventory_sha256': sha(inventory), 'inventory_bytes': len(inventory),
        'measured_task_coverage': None, 'trial_coverage': None, 'selected': None,
        'invalid': None, 'unmeasured': None, 'run_id': None, 'production_authorized': False,
    }
    (args.out / 'summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({k: summary[k] for k in ('rows', 'classes', 'workbooks_opened', 'single_variant_rows',
                                            'default_sides', 'declared_exclusions', 'missing_target_sheets',
                                            'formula_routes', 'eligibility_by_default_side',
                                            'structural_denominator', 'policy_inclusive_candidate_denominator',
                                            'inventory_bytes', 'inventory_sha256')}, sort_keys=True))


if __name__ == '__main__':
    main()
